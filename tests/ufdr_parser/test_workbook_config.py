"""Tests for ufdr_config.toml: loading, overriding, validation, and its effect on the workbook.

Defines:    tests for workbook.config (default, section override, typos, BOM), the init-config
            CLI command, hide_columns, and identifier sources added by configuration.
Used by:    pytest.
Depends on: ufdr_parser (cli, dump, bricks.workbook.config), the test_workbook fixtures.
"""
from __future__ import annotations

import tomllib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.ufdr_parser.test_workbook import REPORT, _bare, _header, _sheet
from ufdr_parser.bricks.workbook.config import default_config_text, load_workbook_config
from ufdr_parser.cli import main
from ufdr_parser.dump import run_dump


def _write(tmp_path: Path, text: str, name: str = "ufdr_config.toml") -> Path:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


@pytest.fixture
def report(tmp_path: Path) -> Path:
    path = tmp_path / "wb_Report.xml"
    path.write_text(REPORT, encoding="utf-8")
    return path


def test_default_is_the_shipped_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    config = load_workbook_config()
    assert config.source == "built-in default"
    assert config.sheets == ("Source device", "Communications", "Contacts", "Identifiers")
    assert [s.model for s in config.identifiers][:2] == ["Party", "PhoneNumber"]
    tomllib.loads(default_config_text())  # the shipped file itself is valid TOML


def test_init_config_writes_the_default_and_never_overwrites(tmp_path: Path) -> None:
    target = tmp_path / "ufdr_config.toml"
    assert main(["init-config", str(target)]) == 0
    assert target.read_text(encoding="utf-8") == default_config_text()
    assert not target.read_bytes().startswith(b"\xef\xbb\xbf")  # TOML forbids a BOM
    target.write_text("# mine", encoding="utf-8")
    assert main(["init-config", str(target)]) == 1
    assert target.read_text(encoding="utf-8") == "# mine"


def test_a_section_replaces_the_default_section(tmp_path: Path) -> None:
    path = _write(tmp_path, '[workbook]\nsheets = ["Identifiers"]\n')
    config = load_workbook_config(path)
    assert config.sheets == ("Identifiers",)
    assert len(config.identifiers) == len(load_workbook_config(_write(tmp_path, "", "empty.toml")).identifiers)
    assert config.source == str(path.resolve())


def test_cwd_file_is_found(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _write(tmp_path, '[workbook]\nsheets = ["Contacts"]\n')
    monkeypatch.chdir(tmp_path)
    assert load_workbook_config().sheets == ("Contacts",)


def test_bom_saved_file_is_accepted(tmp_path: Path) -> None:
    path = tmp_path / "ufdr_config.toml"
    path.write_bytes(b'\xef\xbb\xbf[workbook]\nsheets = ["Contacts"]\n')  # as Windows Notepad saves it
    assert load_workbook_config(path).sheets == ("Contacts",)


@pytest.mark.parametrize(
    "text",
    [
        "[workbok]\nsheets = []\n",                                   # unknown section
        '[workbook]\nsheet = ["Contacts"]\n',                         # unknown key
        '[workbook]\nsheets = "Contacts"\n',                          # not a list
        '[[identifiers]]\nmodel = "Party"\n',                         # missing field
        '[[identifiers]]\nmodel = "Party"\nfield = "x"\nkinds = "y"\n',  # unknown key
    ],
)
def test_typos_are_errors(tmp_path: Path, text: str) -> None:
    with pytest.raises(ValueError):
        load_workbook_config(_write(tmp_path, text))


def test_unknown_sheet_fails_before_the_dump(report: Path, tmp_path: Path) -> None:
    config = _write(tmp_path, '[workbook]\nsheets = ["Contacts", "Contactz"]\n')
    with pytest.raises(ValueError, match="Contactz"):
        run_dump(report, tmp_path / "out", brick_config=config)
    assert not (tmp_path / "out").exists()  # nothing parsed or written


def test_hide_columns(report: Path, tmp_path: Path) -> None:
    config = _write(tmp_path, '[workbook]\nsheets = ["Communications"]\n[hide_columns]\nCommunications = ["Chat uuid", "To basis"]\n')
    run_dump(report, tmp_path / "out", brick_config=config)
    workbook = tmp_path / "out" / "wb_Report_combined.xlsx"
    header = [_bare(c) for c in _header(workbook, "Communications")]
    assert "Chat uuid" not in header and "To basis" not in header and "com_type" in header
    documentation = [row[1] for row in load_workbook(workbook, read_only=True)["Documentation"].iter_rows(values_only=True) if len(row) > 1]
    assert "[report] Chat uuid" not in documentation and "[tool] com_type" in documentation


def test_new_identifier_source_is_one_block(report: Path, tmp_path: Path) -> None:
    config = _write(tmp_path, '''[workbook]
sheets = ["Identifiers"]

[[identifiers]]
model = "Contact"
field = "Name"
kind = "Contact name"
''')
    run_dump(report, tmp_path / "out", brick_config=config)
    rows = _sheet(tmp_path / "out" / "wb_Report_combined.xlsx", "Identifiers")
    assert [(r["Identifier"], r["Kind"], r["CSV file"]) for r in rows] == [("Bob", "Contact name", "wb_Report_Contact.csv")]


def test_contact_child_models_filter(report: Path, tmp_path: Path) -> None:
    config = _write(tmp_path, '[workbook]\nsheets = ["Contacts"]\n[contacts]\nchild_models = ["PhoneNumber"]\n')
    run_dump(report, tmp_path / "out", brick_config=config)
    header = [_bare(c) for c in _header(tmp_path / "out" / "wb_Report_combined.xlsx", "Contacts")]
    assert "PhoneNumber.Value" in header and not any(c.startswith("UserID.") for c in header)


def test_case_name_names_every_output(report: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    summary = run_dump(report, tmp_path / "out", case_name="OP Checkin")
    names = {path.name for path in (tmp_path / "out").iterdir()}
    assert "OP Checkin_combined.xlsx" in names and "OP Checkin_report_info.csv" in names
    assert all(name.startswith("OP Checkin_") for name in names)
    assert all(name.startswith("OP Checkin_") for name in summary.files)
    about = {row[0]: row[1] for row in load_workbook(tmp_path / "out" / "OP Checkin_combined.xlsx", read_only=True)["About"].iter_rows(values_only=True) if len(row) > 1}
    assert about["Case"] == "OP Checkin"
    # Each row still names the report it came from.
    assert _sheet(tmp_path / "out" / "OP Checkin_combined.xlsx", "Communications")[0]["report_file"].endswith("wb_Report.xml")


@pytest.mark.parametrize("bad", ["../escape", "a:b", "   "])
def test_case_name_that_is_not_a_filename_is_refused(report: Path, tmp_path: Path, bad: str) -> None:
    with pytest.raises(ValueError):
        run_dump(report, tmp_path / "out", case_name=bad)
    assert not (tmp_path / "out").exists()
