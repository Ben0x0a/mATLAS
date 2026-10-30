"""Tests for the post-dump brick runner and the combined-workbook brick.

Defines:    a synthetic report covering chats (stated / inferred / not-inferable recipients),
            a standalone message, a call, an email with Cc + attachment, a contact and a
            location, and tests over the resulting <base>_combined.xlsx.
Used by:    pytest.
Depends on: ufdr_parser (dump, bricks), openpyxl.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from ufdr_parser.bricks import base as bricks_base
from ufdr_parser.bricks.workbook import xlsx
from ufdr_parser.dump import run_dump

NS = "http://pa.cellebrite.com/report/2.0"


def _party(pid: str, ident: str, role: str, name: str = "", owner: str = "") -> str:
    name_xml = f'<field name="Name" type="String"><value type="String"><![CDATA[{name}]]></value></field>' if name else ""
    owner_xml = f'<field name="IsPhoneOwner" type="Boolean"><value type="Boolean"><![CDATA[{owner}]]></value></field>' if owner else ""
    ident_xml = f'<field name="Identifier" type="String"><value type="String"><![CDATA[{ident}]]></value></field>' if ident else ""
    return (
        f'<model type="Party" id="{pid}">{ident_xml}'
        f'<field name="Role" type="PartyRole"><value type="PartyRole"><![CDATA[{role}]]></value></field>'
        f"{name_xml}{owner_xml}</model>"
    )


def _user_id(uid: str, category: str, value: str) -> str:
    return (
        f'<model type="UserID" id="{uid}">'
        f'<field name="Category" type="String"><value type="String"><![CDATA[{category}]]></value></field>'
        f'<field name="Value" type="String"><value type="String"><![CDATA[{value}]]></value></field>'
        "</model>"
    )


def _coordinate(cid: str, lat: str, lon: str) -> str:
    return (
        f'<model type="Coordinate" id="{cid}">'
        f'<field name="Latitude" type="Double"><value type="Double"><![CDATA[{lat}]]></value></field>'
        f'<field name="Longitude" type="Double"><value type="Double"><![CDATA[{lon}]]></value></field>'
        "</model>"
    )


def _message(mid: str, ts: str, body: str, parties: str) -> str:
    return (
        f'<model type="InstantMessage" id="{mid}">'
        f'<field name="TimeStamp" type="TimeStamp"><value type="TimeStamp">{ts}</value></field>'
        f'<field name="Body" type="String"><value type="String"><![CDATA[{body}]]></value></field>'
        f'<multiModelField name="From" type="Party">{parties}</multiModelField>'
        "</model>"
    )


REPORT = f"""<?xml version="1.0" encoding="utf-8"?>
<project xmlns="{NS}" name="wb" reportVersion="8.5" extractionType="FileSystem">
  <sourceExtractions>
    <extractionInfo id="0" name="File System" type="FileSystem" deviceName="IPHONE" index="0" />
  </sourceExtractions>
  <caseInformation>
    <field name="Case name" fieldType="CaseName">Test case</field>
  </caseInformation>
  <metadata section="Extraction Data">
    <item name="DeviceInfoExtractionStartDateTime" sourceExtraction="0"><![CDATA[03/09/2025 17:35:59(UTC+2)]]></item>
  </metadata>
  <images>
    <image key="FileDump" path="EXTRACTION_FFS.zip" size="10">
      <metadata section="Hashes"><item name="SHA256"><![CDATA[ABCD]]></item></metadata>
    </image>
  </images>
  <taggedFiles />
  <decodedData>
    <modelType type="Chat">
      <model type="Chat" id="chat1">
        <field name="Source" type="String"><value type="String"><![CDATA[WhatsApp]]></value></field>
        <field name="Name" type="String"><value type="String"><![CDATA[Famille]]></value></field>
        <multiModelField name="Participants" type="Party">
          {_party("cp1", "+41790000001", "General", "Alice", "True")}
          {_party("cp2", "+41790000002", "General", "Bob")}
        </multiModelField>
        <multiModelField name="Messages" type="InstantMessage">
          {_message("m2", "2025-01-01T10:05:00.000+00:00", "=SUM(1,2)", _party("p2", "+41790000002", "From", "Bob") + _party("p3", "+41790000001", "To", "Alice", "True"))}
          {_message("m1", "2025-01-01T10:00:00.000+00:00", "Salut, ça va ?", _party("p1", "+41790000001", "From", "Alice", "True") + _party("pe", "", "To"))}
          {_message("m3", "2025-01-01T10:10:00.000+00:00", "stranger", _party("p4", "+41799999999", "From"))}
        </multiModelField>
      </model>
    </modelType>
    <modelType type="InstantMessage">
      {_message("im1", "2025-01-02T08:00:00.000+00:00", "SMS seul", _party("p5", "+41790000003", "From"))}
    </modelType>
    <modelType type="Call">
      <model type="Call" id="call1">
        <field name="Direction" type="String"><value type="String"><![CDATA[Outgoing]]></value></field>
        <field name="TimeStamp" type="TimeStamp"><value type="TimeStamp">2025-01-03T09:00:00.000+00:00</value></field>
        <field name="Duration" type="TimeSpan"><value type="TimeSpan"><![CDATA[00:01:00]]></value></field>
        <multiModelField name="Parties" type="Party">
          {_party("c1", "+41790000001", "From", "Alice", "True")}
          {_party("c2", "+41790000002", "To", "Bob")}
        </multiModelField>
      </model>
    </modelType>
    <modelType type="Email">
      <model type="Email" id="mail1">
        <field name="TimeStamp" type="TimeStamp"><value type="TimeStamp">2025-01-04T12:00:00.000</value></field>
        <field name="Subject" type="String"><value type="String"><![CDATA[Réunion]]></value></field>
        <multiModelField name="Parties" type="Party">
          {_party("e1", "a@x.ch", "From")}
          {_party("e2", "b@x.ch", "To")}
          {_party("e3", "c@x.ch", "Cc")}
        </multiModelField>
        <multiModelField name="Attachments" type="Attachment">
          <model type="Attachment" id="att1">
            <field name="Filename" type="String"><value type="String"><![CDATA[plan.pdf]]></value></field>
            <field name="attachment_extracted_path" type="String"><value type="String"><![CDATA[files/plan.pdf]]></value></field>
          </model>
        </multiModelField>
      </model>
    </modelType>
    <modelType type="Contact">
      <model type="Contact" id="ct1">
        <field name="Source" type="String"><value type="String"><![CDATA[WhatsApp]]></value></field>
        <field name="Name" type="String"><value type="String"><![CDATA[Bob]]></value></field>
        <multiModelField name="Entries" type="PhoneNumber">
          <model type="PhoneNumber" id="ph1">
            <field name="Category" type="String"><value type="String"><![CDATA[Mobile]]></value></field>
            <field name="Value" type="String"><value type="String"><![CDATA[+41790000002]]></value></field>
          </model>
        </multiModelField>
        <multiModelField name="Ids" type="UserID">
          {_user_id("u1", "WhatsApp User Id", "41790000002@s.whatsapp.net")}
          {_user_id("u2", "Push Name", "Bobby")}
          {_user_id("u3", "Instagram Id", "bob.insta")}
        </multiModelField>
      </model>
      <model type="Contact" id="ct2">
        <field name="Source" type="String"><value type="String"><![CDATA[InteractionC]]></value></field>
        <multiModelField name="Ids" type="UserID">
          {_user_id("u4", "Identifier", "abc-123")}
        </multiModelField>
      </model>
    </modelType>
    <modelType type="Location">
      <model type="Location" id="loc1">
        <field name="TimeStamp" type="TimeStamp"><value type="TimeStamp">2025-01-05T07:00:00.000+00:00</value></field>
        <modelField name="Position" type="Coordinate">
          <model type="Coordinate" id="co1">
            <field name="Latitude" type="Double"><value type="Double"><![CDATA[46.5]]></value></field>
            <field name="Longitude" type="Double"><value type="Double"><![CDATA[6.6]]></value></field>
          </model>
        </modelField>
      </model>
    </modelType>
    <modelType type="DeviceInfoEntry">
      <model type="DeviceInfoEntry" id="dev1">
        <field name="EntryName" type="String"><value type="String"><![CDATA[IMEI]]></value></field>
        <field name="EntryValue" type="String"><value type="String"><![CDATA[353872860051601]]></value></field>
        <field name="EntryCategory" type="String"><value type="String"><![CDATA[General]]></value></field>
        <field name="Source" type="String"><value type="String"><![CDATA[Native]]></value></field>
      </model>
    </modelType>
    <modelType type="CellTower">
      <model type="CellTower" id="cell1">
        <field name="TimeStamp" type="TimeStamp"><value type="TimeStamp">2025-01-05T06:00:00.000+00:00</value></field>
        <field name="MCC" type="String"><value type="String"><![CDATA[228]]></value></field>
        <field name="CID" type="String"><value type="String"><![CDATA[42]]></value></field>
        <modelField name="Position" type="Coordinate">{_coordinate("cc1", "46.1", "6.1")}</modelField>
      </model>
    </modelType>
    <modelType type="WirelessNetwork">
      <model type="WirelessNetwork" id="wifi1">
        <field name="SSId" type="String"><value type="String"><![CDATA[Home]]></value></field>
        <field name="LastConnection" type="TimeStamp"><value type="TimeStamp">2025-01-05T08:00:00.000</value></field>
      </model>
    </modelType>
    <modelType type="Journey">
      <model type="Journey" id="j1">
        <field name="Name" type="String"><value type="String"><![CDATA[To the lake]]></value></field>
        <field name="StartTime" type="TimeStamp"><value type="TimeStamp">2025-01-05T09:00:00.000+00:00</value></field>
        <multiModelField name="WayPoints" type="Location">
          <model type="Location" id="wp1">
            <field name="Description" type="String"><value type="String"><![CDATA[Turn left]]></value></field>
            <modelField name="Position" type="Coordinate">{_coordinate("wc1", "46.2", "6.2")}</modelField>
          </model>
          <model type="Location" id="wp2">
            <field name="TimeStamp" type="TimeStamp"><value type="TimeStamp">2025-01-05T09:30:00.000+00:00</value></field>
            <modelField name="Position" type="Coordinate">{_coordinate("wc2", "46.3", "6.3")}</modelField>
          </model>
        </multiModelField>
      </model>
    </modelType>
    <modelType type="SearchedItem">
      <model type="SearchedItem" id="s1">
        <field name="TimeStamp" type="TimeStamp"><value type="TimeStamp">2025-01-05T10:00:00.000+00:00</value></field>
        <field name="Value" type="String"><value type="String"><![CDATA[lake geneva]]></value></field>
      </model>
    </modelType>
    <modelType type="SocialMediaActivity">
      <model type="SocialMediaActivity" id="sm1">
        <field name="TimeStamp" type="TimeStamp"><value type="TimeStamp">2025-01-05T11:00:00.000+00:00</value></field>
        <modelField name="Position" type="Coordinate">{_coordinate("smc1", "46.4", "6.4")}</modelField>
      </model>
      <model type="SocialMediaActivity" id="sm2">
        <field name="TimeStamp" type="TimeStamp"><value type="TimeStamp">2025-01-05T12:00:00.000+00:00</value></field>
      </model>
    </modelType>
  </decodedData>
  <extraInfos>
    <extraInfo type="model" id="chat1"><sourceInfo><nodeInfos><nodeInfo name="msgstore.db" path="FFS.zip/msgstore.db" tableName="messages" offset="10" /></nodeInfos></sourceInfo></extraInfo>
  </extraInfos>
</project>
"""


def _header(path: Path, name: str) -> tuple[str, ...]:
    return next(load_workbook(path, read_only=True)[name].iter_rows(values_only=True))


def _bare(column: str) -> str:
    for prefix in ("[report] ", "[tool] "):
        if column.startswith(prefix):
            return column[len(prefix):]
    return column


def _sheet(path: Path, name: str) -> list[dict[str, str | None]]:
    """Rows keyed by BARE column names (origin prefixes are asserted by test_column_origins)."""
    workbook = load_workbook(path, read_only=True)
    rows = workbook[name].iter_rows(values_only=True)
    header = [_bare(column) for column in next(rows)]
    return [dict(zip(header, row)) for row in rows]


@pytest.fixture
def report(tmp_path: Path) -> Path:
    path = tmp_path / "wb_Report.xml"
    path.write_text(REPORT, encoding="utf-8")
    return path


@pytest.fixture
def workbook(report: Path, tmp_path: Path) -> Path:
    summary = run_dump(report, tmp_path / "out")
    assert [(b.name, b.status) for b in summary.bricks] == [("workbook", "ok")]
    path = tmp_path / "out" / "wb_Report_combined.xlsx"
    assert summary.bricks[0].files == [path]
    return path


def test_sheets_in_order(workbook: Path) -> None:
    assert load_workbook(workbook, read_only=True).sheetnames == [
        "About", "Documentation", "Source device", "Communications", "Messages", "Calls",
        "Emails", "Contacts", "Locations", "Identifiers", "CDRImporter",
    ]


def test_messages_grouped_by_chat_sorted_by_time_with_to_basis(workbook: Path) -> None:
    rows = _sheet(workbook, "Messages")
    assert [r["Record uuid"] for r in rows] == ["m1", "m2", "m3", "im1"]
    m1, m2, m3, im1 = rows
    # m1 states only an EMPTY To party (ignored), so To is inferred from the participants.
    assert m1["From"] == "Alice <+41790000001> [flagged as phone owner]"
    assert m1["To"] == "Bob <+41790000002>"
    assert m1["To basis"] == "inferred from chat participants"
    assert m1["Body"] == "Salut, ça va ?"
    assert m1["Chat"] == "Famille" and m1["Chat app"] == "WhatsApp"
    assert m2["To basis"] == "stated"
    # The sender is not a participant: "everyone else" cannot be told, so nothing is inferred.
    assert m3["To"] is None and m3["To basis"] == "not stated"
    assert im1["Category"] == "Instant message" and im1["To basis"] == "not stated"
    # Nested rows inherit the chat's source; every row names the report.
    assert m1["source_path"] == "FFS.zip/msgstore.db"
    assert m1["report_file"].endswith("wb_Report.xml")


def test_values_stay_text(workbook: Path) -> None:
    cells = {c.value: c.data_type for row in load_workbook(workbook)["Messages"].iter_rows() for c in row}
    assert cells["=SUM(1,2)"] == "s"  # not a live formula
    rows = _sheet(workbook, "CDRImporter")
    assert rows[0]["A_MSISDN"] == "+41790000001"  # leading + kept


def test_calls_emails_and_unified(workbook: Path) -> None:
    call = _sheet(workbook, "Calls")[0]
    assert (call["Direction"], call["Duration"], call["To"]) == ("Outgoing", "00:01:00", "Bob <+41790000002>")
    email = _sheet(workbook, "Emails")[0]
    assert email["Other parties"] == "[Cc] c@x.ch"
    assert email["Attachments"] == "plan.pdf (files/plan.pdf)"
    assert email["TimeStamp.offset"] is None  # no offset determinable: blank, no default
    unified = _sheet(workbook, "Communications")
    assert [r["Category"] for r in unified] == [
        "Chat message", "Chat message", "Chat message", "Instant message", "Call", "Email",
    ]


def test_cdr_one_row_per_event(workbook: Path) -> None:
    rows = _sheet(workbook, "CDRImporter")
    ids = [row["Record uuid"] for row in rows]
    assert len(ids) == len(set(ids)) == len(_sheet(workbook, "Communications"))  # one row per event
    by_record = {row["Record uuid"]: row for row in rows}
    assert by_record["m1"]["B_MSISDN"] == "+41790000002"
    assert by_record["m1"]["B_basis"] == "inferred from chat participants"
    assert by_record["m3"]["B_MSISDN"] is None and by_record["m3"]["B_basis"] == "not stated"
    assert by_record["call1"]["Type2"] == "Outgoing"
    assert by_record["mail1"]["B_MSISDN"] == "b@x.ch"  # Cc is not a recipient in B


def test_identifiers_contacts_locations(workbook: Path) -> None:
    ids = {r["Identifier"]: r for r in _sheet(workbook, "Identifiers")}
    alice = ids["+41790000001"]
    assert alice["Flagged as phone owner"] == "yes"
    assert alice["As sender"] == "2" and alice["As recipient"] == "1"
    assert "Chat (participant) x1" in alice["Seen in"]
    contact = _sheet(workbook, "Contacts")[0]
    assert contact["Phone numbers"] == "Mobile: +41790000002"
    location = next(r for r in _sheet(workbook, "Locations") if r["Record uuid"] == "loc1")
    assert (location["Latitude"], location["Longitude"], location["TimeStamp.offset"]) == ("46.5", "6.6", "+00:00")


def test_about_documents_every_sheet(workbook: Path) -> None:
    about = [row for row in load_workbook(workbook, read_only=True)["About"].iter_rows(values_only=True)]
    sheets = {row[0]: row[2] for row in about if row and row[0] in ("Messages", "Calls", "CDRImporter") and row[2] and row[2].isdigit()}
    assert sheets == {"Messages": "4", "Calls": "1", "CDRImporter": "6"}
    documentation = list(load_workbook(workbook, read_only=True)["Documentation"].iter_rows(values_only=True))
    assert any(row[:3] == ("Messages", "[tool] To basis", "tool") for row in documentation)
    assert any(row[:3] == ("Calls", "[report] Direction", "report") for row in documentation)


def test_parallel_parse_gives_same_workbook(report: Path, tmp_path: Path, workbook: Path) -> None:
    run_dump(report, tmp_path / "par", parse_workers=2)
    parallel = tmp_path / "par" / "wb_Report_combined.xlsx"
    for name in ("Messages", "Communications", "CDRImporter", "Identifiers"):
        assert _sheet(parallel, name) == _sheet(workbook, name)


def test_no_bricks_writes_csvs_only(report: Path, tmp_path: Path) -> None:
    summary = run_dump(report, tmp_path / "out", bricks=set())
    assert summary.bricks == []
    assert not list((tmp_path / "out").glob("*.xlsx"))
    assert summary.files


def test_failing_brick_does_not_fail_dump(report: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    class Broken:
        name = "broken"
        description = "always fails"

        def run(self, output):
            raise RuntimeError("boom")

    monkeypatch.setattr(bricks_base, "_REGISTRY", {"broken": Broken()})
    summary = run_dump(report, tmp_path / "out")
    assert summary.bricks[0].status == "failed" and "boom" in summary.bricks[0].error
    assert summary.files  # the CSVs stand


def test_excel_limits(report: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(xlsx, "EXCEL_MAX_ROWS", 3)  # header + 2 rows per sheet part
    monkeypatch.setattr(xlsx, "EXCEL_MAX_CELL_CHARS", 160)
    path = tmp_path / "limits.xml"
    path.write_text(REPORT.replace("SMS seul", "x" * 500), encoding="utf-8")
    run_dump(path, tmp_path / "out")
    book = load_workbook(tmp_path / "out" / "limits_combined.xlsx", read_only=True)
    assert "Messages (2)" in book.sheetnames
    rows = _sheet(tmp_path / "out" / "limits_combined.xlsx", "Messages (2)")
    body = rows[-1]["Body"]
    assert len(body) == 160 and "(500 characters)" in body


def test_control_characters_are_escaped_not_dropped() -> None:
    # XML 1.0 cannot carry U+000B, so it cannot come through a report; CSV text can.
    stats = xlsx.WriteStats()
    sheet = xlsx.open_workbook().create_sheet("t")
    cell = xlsx.text_cell(sheet, "a\x0bb", stats)
    assert cell.value == "a[U+000B]b" and stats.escaped_characters == 1


def test_contact_identifier_columns(workbook: Path) -> None:
    bob, other = _sheet(workbook, "Contacts")
    # Category names the app first; else the contact's Source (WhatsApp) claims the entry.
    assert bob["WhatsApp"] == "WhatsApp User Id: 41790000002@s.whatsapp.net; Push Name: Bobby"
    assert bob["Instagram"] == "Instagram Id: bob.insta"
    assert bob["Phone numbers"] == "Mobile: +41790000002"
    # Every configured column exists, blank when the contact has none.
    assert "Facebook" in bob and bob["Facebook"] is None
    assert bob["Other identifiers"] is None
    assert other["Other identifiers"] == "Identifier: abc-123"


@pytest.mark.parametrize(
    ("category", "source", "expected"),
    [
        ("Facebook Id", None, "Facebook"),
        ("Username", "Facebook Messenger", "Facebook"),
        ("User ID", "Snapchat", "Snapchat"),
        ("Identifier", "Native Messages", "Other identifiers"),
        ("Id", "X", "X (Twitter)"),
        ("Index", "Xperia", "Other identifiers"),  # whole words only
    ],
)
def test_classify_identifier(category: str, source: str | None, expected: str) -> None:
    from ufdr_parser.bricks.workbook.sheets import classify_identifier

    assert classify_identifier(category, source) == expected


def test_column_origins(workbook: Path) -> None:
    messages = _header(workbook, "Messages")
    assert messages[:3] == ("[tool] Category", "[report] Chat", "[report] Chat type")
    assert "[tool] TimeStamp.offset" in messages and "[report] TimeStamp" in messages
    assert "[report] source_path" in messages and "[tool] report_file" in messages
    cdr = _header(workbook, "CDRImporter")
    assert cdr[0] == "A_MSISDN"  # legacy importer names stay exact
    assert "[tool] DateTime offset" in cdr and "[report] Source" in cdr
    for sheet in ("Communications", "Messages", "Calls", "Emails", "CDRImporter"):
        assert all(_bare(column) != "App" for column in _header(workbook, sheet))


def test_source_device(workbook: Path) -> None:
    rows = _sheet(workbook, "Source device")
    by_name = {(r["Section"], r["Name"]): r for r in rows}
    assert by_name[("Case information", "Case name")]["Value"] == "Test case"
    start = by_name[("Metadata: Extraction Data", "DeviceInfoExtractionStartDateTime")]
    assert start["Value"] == "03/09/2025 17:35:59(UTC+2)" and start["Details"] == "sourceExtraction=0"
    assert by_name[("Image FileDump", "SHA256")]["Value"] == "ABCD"
    imei = by_name[("Device info: General", "IMEI")]
    assert imei["Value"] == "353872860051601" and imei["Record uuid"] == "dev1"
    assert [r["Section"] for r in rows][:1] == ["Report"]  # header first, decoded device info last


def test_locations_all_position_types(workbook: Path) -> None:
    rows = _sheet(workbook, "Locations")
    header = _header(workbook, "Locations")
    for removed in ("Time", "Time field", "Label", "Address", "App"):
        assert all(_bare(column) != removed for column in header)
    # Report order, type by type, no sorting.
    assert [r["Record uuid"] for r in rows] == ["loc1", "cell1", "wifi1", "wp1", "wp2", "s1", "sm1"]
    by_id = {r["Record uuid"]: r for r in rows}
    cell = by_id["cell1"]
    assert (cell["Record type"], cell["Latitude"], cell["Latitude field"]) == ("CellTower", "46.1", "SMCoordinate.Latitude")
    assert (cell["MCC"], cell["CID"]) == ("228", "42")  # own fields verbatim, not merged into a label
    wifi = by_id["wifi1"]
    assert (wifi["Latitude"], wifi["Latitude field"], wifi["LastConnection.offset"]) == (None, None, None)
    wp1 = by_id["wp1"]
    assert wp1["Record type"] == "Journey point" and wp1["Description"] == "Turn left"
    assert (wp1["Latitude"], wp1["Longitude field"]) == ("46.2", "SSMCoordinate.Longitude")
    assert (wp1["Journey.Name"], wp1["Journey.StartTime"]) == ("To the lake", "2025-01-05T09:00:00.000+00:00")
    assert wp1["TimeStamp"] is None  # never filled from the journey
    assert by_id["s1"]["Value"] == "lake geneva" and by_id["s1"]["Latitude"] is None
    # An unconfigured type is listed only for its records that carry a coordinate.
    assert "sm2" not in by_id


def test_cdr_group_message_is_one_row(tmp_path: Path) -> None:
    """A chat message with several recipients gives ONE CDR row naming the group, not one per
    member (per-member rows multiplied group messages by the group size)."""
    members = "".join(_party(f"gp{i}", f"+4179000010{i}", "General") for i in range(4))
    group = f"""<?xml version="1.0" encoding="utf-8"?>
<project xmlns="{NS}"><decodedData><modelType type="Chat">
  <model type="Chat" id="g1">
    <field name="Id" type="String"><value type="String"><![CDATA[120363@g.us]]></value></field>
    <field name="Name" type="String"><value type="String"><![CDATA[Team]]></value></field>
    <multiModelField name="Participants" type="Party">{members}</multiModelField>
    <multiModelField name="Messages" type="InstantMessage">
      {_message("gm1", "2025-02-01T10:00:00.000+00:00", "inferred", _party("gs1", "+41790000100", "From"))}
      {_message("gm2", "2025-02-01T10:01:00.000+00:00", "stated", _party("gs2", "+41790000101", "From") + _party("gt1", "+41790000102", "To") + _party("gt2", "+41790000103", "To"))}
    </multiModelField>
  </model>
</modelType></decodedData><extraInfos /></project>
"""
    path = tmp_path / "group.xml"
    path.write_text(group, encoding="utf-8")
    run_dump(path, tmp_path / "out")
    rows = _sheet(tmp_path / "out" / "group_combined.xlsx", "CDRImporter")
    assert [(r["Record uuid"], r["B_MSISDN"]) for r in rows] == [("gm1", "120363@g.us"), ("gm2", "120363@g.us")]
    assert rows[0]["B_basis"].startswith("inferred from chat participants; group chat (3 recipients")
    assert rows[1]["B_basis"].startswith("stated; group chat (2 recipients")
    # The members themselves stay on the Messages sheet.
    message = next(r for r in _sheet(tmp_path / "out" / "group_combined.xlsx", "Messages") if r["Record uuid"] == "gm1")
    assert message["To"].count("+4179000010") == 3
