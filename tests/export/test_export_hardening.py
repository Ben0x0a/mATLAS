"""Security/integrity-hardening tests for the export bricks (FIX-17).

Covers three defects the review found in the text/spreadsheet exporters, exercised
end-to-end over the hand-authored manual store (tests/export/_pipeline.build_manual_store),
with a few targeted SQL mutations to plant the hostile/edge values the golden fixtures
cannot carry:

  1. XLSX Int64 precision: an ``s2_cell_id`` = int64-max survives the XLSX round-trip
     byte-exact (the old ``DataFrame.to_excel`` floated it beyond 2**53).
  2. Formula injection: a ``=HYPERLINK(...)`` cell is written INERT (never a live formula)
     in the XLSX and BOTH CSV writers, while re-importing to the exact original string.
  3. OverflowError guard: one out-of-range int64 epoch yields an empty time cell for that
     row instead of aborting the whole .kml / GeoVisualizer file.
"""
from __future__ import annotations

import csv
from pathlib import Path
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from model_atlas.core import vocabularies as V
from model_atlas.core.flat_view import build_flat_view
from model_atlas.core.store import RunStore
from model_atlas.export._shared import FORMULA_GUARD
from model_atlas.export.csv_flat import CsvFlatExporter, ExportContext
from model_atlas.export.geovisualizer import GeovisualizerExporter
from model_atlas.export.kml import KmlExporter
from model_atlas.export.xlsx import XlsxExporter

from tests.export._pipeline import build_manual_store

_INT64_MAX = 9223372036854775807   # 2**63 - 1: an s2_cell_id far beyond float64's exact range
_FORMULA_VALUE = '=HYPERLINK("http://evil.example/steal","click me")'
_KML_NS = "{http://www.opengis.net/kml/2.2}"


# --- 1. XLSX Int64 precision --------------------------------------------------

def test_xlsx_preserves_large_s2_cell_id_exactly(tmp_path: Path) -> None:
    with RunStore() as store:
        build_manual_store(store)
        # Plant an int64-max S2 cell id on the instant observation's spatial integration.
        store.write_spatial_index(V.SpatialIndex(
            spatial_index_id="sidx-instant", spatial_integration_id="si-instant",
            s2_cell_id=_INT64_MAX, s2_level=30))
        frame = build_flat_view(store.connection, include_source_columns=False)
        assert frame["s2_cell_id_matlas"].dtype == "Int64"   # nullable-Int64 (the corruption route)
        written = XlsxExporter().export(
            store, ExportContext(output_dir=tmp_path, base_name="merged"))

    worksheet = load_workbook(written[0]).active
    header = [cell.value for cell in worksheet[1]]
    column_index = header.index("s2_cell_id_matlas") + 1
    planted = [
        worksheet.cell(row=row, column=column_index).value
        for row in range(2, worksheet.max_row + 1)
        if worksheet.cell(row=row, column=column_index).value not in (None, "")
    ]
    # The exact int64-max survives (written as text, so openpyxl never floats it): the old
    # to_excel path stored 9.223372036854776e+18 — a WRONG cell id — here.
    assert planted == [str(_INT64_MAX)]
    assert int(planted[0]) == _INT64_MAX


# --- 2. formula injection neutralised in xlsx + both CSVs ---------------------

def _mutate_formula_entity(store: RunStore) -> None:
    """Make the instant observation's entity/concerns a live-formula string.

    ``concerns`` (entity.concerns) is the GeoVisualizer ``Target`` and a flat-CSV
    column; ``entity`` (entity.identifier) is another flat-CSV column — so one mutation
    plants the hostile value into a cell of every text export under test.
    """
    store.connection.execute(
        "UPDATE entity SET identifier = ?, concerns = ? WHERE entity_id = 'entity-instant'",
        (_FORMULA_VALUE, _FORMULA_VALUE))


def test_csv_flat_neutralises_formula_and_round_trips(tmp_path: Path) -> None:
    with RunStore() as store:
        build_manual_store(store)
        _mutate_formula_entity(store)
        written = CsvFlatExporter().export(
            store, ExportContext(output_dir=tmp_path, base_name="merged",
                                 include_source_columns=False))

    records = _reread(written[0])
    cells = ([row["entity"] for row in records]
             + [row["raw_identifier"] for row in records])
    guarded = [value for value in cells if value.startswith(FORMULA_GUARD)]
    assert guarded, "the formula cell should be apostrophe-guarded"
    for value in guarded:
        assert not value.startswith("=")            # never a leading formula trigger
        assert value[len(FORMULA_GUARD):] == _FORMULA_VALUE   # strips to the exact original


def test_geovisualizer_neutralises_formula(tmp_path: Path) -> None:
    with RunStore() as store:
        build_manual_store(store)
        _mutate_formula_entity(store)
        written = GeovisualizerExporter().export(
            store, ExportContext(output_dir=tmp_path, base_name="merged"))

    targets = [row["Target"] for row in _reread(written[0])]
    guarded = [value for value in targets if value.startswith(FORMULA_GUARD)]
    assert guarded, "the Target formula cell should be apostrophe-guarded"
    for value in guarded:
        assert not value.startswith("=")
        assert value[len(FORMULA_GUARD):] == _FORMULA_VALUE


def test_xlsx_writes_formula_as_inert_text(tmp_path: Path) -> None:
    with RunStore() as store:
        build_manual_store(store)
        _mutate_formula_entity(store)
        written = XlsxExporter().export(
            store, ExportContext(output_dir=tmp_path, base_name="merged"))

    worksheet = load_workbook(written[0]).active
    header = [cell.value for cell in worksheet[1]]
    column_index = header.index("entity") + 1
    formula_cells = [
        worksheet.cell(row=row, column=column_index)
        for row in range(2, worksheet.max_row + 1)
        if worksheet.cell(row=row, column=column_index).value == _FORMULA_VALUE
    ]
    assert formula_cells, "the formula string should be present as a value"
    for cell in formula_cells:
        # Data-type 's' (string), NOT 'f' (formula): inert on open, yet byte-exact so it
        # re-imports to the original string.
        assert cell.data_type == "s"
        assert cell.value == _FORMULA_VALUE


# --- 3. out-of-range epoch skips one feature, never the file ------------------

def _plant_bad_epoch(store: RunStore) -> None:
    """Corrupt the instant observation's lower/upper epoch to an int64-max micro value —
    a datetime-range overflow when localised (the raiser Part 3 guards)."""
    store.connection.execute(
        "UPDATE timespan_integration SET lower_unix_utc_us = ?, upper_unix_utc_us = ? "
        "WHERE timespan_integration_id = 'ti-instant'",
        (_INT64_MAX, _INT64_MAX))


def test_kml_survives_out_of_range_epoch(tmp_path: Path) -> None:
    with RunStore() as store:
        build_manual_store(store)
        _plant_bad_epoch(store)
        written = KmlExporter().export(
            store, ExportContext(output_dir=tmp_path, base_name="merged"))

    root = ET.fromstring(written[0].read_bytes())
    placemarks = root.findall(f".//{_KML_NS}Placemark")
    # BOTH located rows are still exported (the file is not aborted); the corrupt-epoch row
    # simply carries no TimeStamp.
    assert len(placemarks) == 2
    without_time = [
        placemark for placemark in placemarks
        if placemark.find(f"{_KML_NS}TimeStamp/{_KML_NS}when") is None
    ]
    assert len(without_time) == 1


def test_geovisualizer_survives_out_of_range_epoch(tmp_path: Path) -> None:
    with RunStore() as store:
        build_manual_store(store)
        _plant_bad_epoch(store)
        written = GeovisualizerExporter().export(
            store, ExportContext(output_dir=tmp_path, base_name="merged"))

    out = _reread(written[0])
    # Instant (1 row) + interval (2 bound rows) = 3 rows, none lost to the bad epoch.
    assert len(out) == 3
    blank_time = [row for row in out if row["Timestamp (local)"] == ""]
    assert len(blank_time) == 1


# --- small helper -------------------------------------------------------------

def _reread(path: Path) -> list[dict]:
    """Every data row of a ';'-delimited utf-8-sig CSV as a dict (verbatim cell text)."""
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=";"))
